"""Tiny generated documents for Library extraction tests."""

from __future__ import annotations

from pathlib import Path

from docx import Document


def write_docx(path: Path, phrase: str) -> None:
    document = Document()
    document.add_paragraph(phrase)
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = phrase + " table"
    document.save(path)


def write_pdf(path: Path, phrase: str) -> None:
    stream = f"BT /F1 12 Tf 72 720 Td ({phrase}) Tj ET".encode("ascii")
    objects = (
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /Resources << /Font << /F1 5 0 R >> >> /MediaBox [0 0 612 792] /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    )
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, content in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode())
        output.extend(content)
        output.extend(b"\nendobj\n")
    startxref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    output.extend(b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:]))
    output.extend(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{startxref}\n%%EOF\n".encode())
    path.write_bytes(output)
