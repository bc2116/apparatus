"""Format-specific text extraction for Library source files."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from io import BytesIO
from contextlib import redirect_stderr
import io

from docx import Document
from pypdf import PdfReader

EXTRACTOR_VERSION = "1"


@dataclass(frozen=True)
class Extraction:
    status: str
    extractor: str
    text: str = ""
    error: str | None = None


def extract_bytes(path: Path, content: bytes) -> Extraction:
    """Extract text from one supported source without treating it as instructions."""
    suffix = path.suffix.lower()
    try:
        if suffix in {".md", ".txt"}:
            text = content.decode("utf-8", errors="replace")
            return _result(text, "utf-8")
        if suffix == ".docx":
            document = Document(BytesIO(content))
            parts = [paragraph.text for paragraph in document.paragraphs]
            parts.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
            return _result("\n".join(parts), "python-docx")
        if suffix == ".pdf":
            with redirect_stderr(io.StringIO()):
                reader = PdfReader(BytesIO(content))
                return _result("\n".join(page.extract_text() or "" for page in reader.pages), "pypdf")
    except Exception as error:
        return Extraction("error", {".docx": "python-docx", ".pdf": "pypdf"}.get(suffix, "utf-8"), error="source could not be extracted")
    return Extraction("unsupported", "none", error=f"unsupported file type: {suffix or 'none'}")


def _result(text: str, extractor: str) -> Extraction:
    return Extraction("extracted" if text else "no_text", extractor, text)
